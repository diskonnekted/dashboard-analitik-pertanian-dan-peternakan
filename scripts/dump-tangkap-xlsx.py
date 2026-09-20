import openpyxl, glob, os

base = r"I:\pertanian\pertanian-2\public\14. Distankan KP\Produksi dan Nilai Produksi Perikanan Tangkap Menurut Kecamatan dan Jenis Penangkapan\_tmp"
files = sorted(glob.glob(os.path.join(base, "*.xlsx")))
print("Files:", [os.path.basename(f) for f in files])
for f in files:
    if "2019" not in os.path.basename(f) and "2024" not in os.path.basename(f):
        continue
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f"\n=== {os.path.basename(f)} | sheet: {ws.title} | dims: {ws.max_row}x{ws.max_column} ===")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            vals = [str(v)[:28] if v is not None else "" for v in row]
            print(f"  r{i+1}: {' | '.join(vals)}")
            if i >= 13:
                print("  ...")
                break
    wb.close()
