import openpyxl, os

base = r"I:\pertanian\pertanian-2\public\14. Distankan KP\Produksi dan Nilai Produksi Perikanan Tangkap Menurut Kecamatan dan Jenis Penangkapan\_tmp"
for yr in ("2019", "2024"):
    f = os.path.join(base, f"Produksi dan Nilai Produksi Perikanan Tangkap Menurut Kecamatan dan Jenis Penangkapan {yr}.xlsx")
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    print(f"===== {yr} (dims {ws.max_row}x{ws.max_column}) =====")
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = [("" if v is None else str(v)[:24]) for v in row]
        if any(vals):
            print(f"r{i:03d}: " + " | ".join(f"[{j}]{v}" for j, v in enumerate(vals) if v != ""))
    wb.close()
    print()
